from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / ".tmp_shrink_paper"
OUT_DIR.mkdir(exist_ok=True)
OUT = OUT_DIR / "07_thesis_school_recommendation_short_paper.docx"


TITLE_CN = "基于大语言模型的缺陷检测AGENT设计与实现"
TITLE_EN = "Design and Implementation of a Defect Detection Agent Based on a Large Language Model"
AUTHOR_LINE = "人工智能与交通工程学院 信管（闽台）2202班：颜益炜       指导教师：郑积仕"


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc: Document, text: str, *, size=12, bold=False, align=None, first_line=True, font="宋体"):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(20)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if first_line:
        p.paragraph_format.first_line_indent = Pt(24)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, font, size, bold)
    return p


def add_mixed_label_para(doc: Document, label: str, body: str, *, english=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(20 if not english else 16)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run(label)
    set_run_font(r, "宋体" if not english else "Times New Roman", 12, True)
    r = p.add_run(body)
    set_run_font(r, "宋体" if not english else "Times New Roman", 12, False)
    return p


def add_section(doc: Document, title: str):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(20)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run(title)
    set_run_font(r, "宋体", 12, True)
    return p


def add_ref(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(16)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Pt(0)
    r = p.add_run(text)
    set_run_font(r, "宋体", 10.5, False)
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = Pt(20)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = Pt(24)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(TITLE_CN)
    set_run_font(r, "宋体", 16, False)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = Pt(20)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(AUTHOR_LINE)
    set_run_font(r, "楷体_GB2312", 12, False)

    add_mixed_label_para(
        doc,
        "摘要：",
        "DAC-3D检测软件在实际使用中涉及参数设置、扫描区域确认、运行状态查看、检测结果判读和异常处理等多个环节。传统界面能够支撑基础检测控制，但在知识检索、自然语言操作、结果解释、操作确认和过程追踪方面仍存在一定不足。围绕上述问题，本文设计并实现了一套面向DAC-3D场景的大语言模型智能交互助手。系统采用前后端分离与模块化设计，以统一Agent运行时为入口，结合知识库检索、上下文工程、技能选择、会话记忆、工具网关、安全审查、DAC-3D适配、结果解析和Trace/Eval等模块完成请求处理。系统不改变DAC-3D主检测算法和主检测流程，而是在主系统之外构建受知识约束、受工具网关控制并且可追踪验证的辅助交互层。测试结果表明，该方案能够在本地原型环境中支撑参数问答、命令预览、状态查询、结果解释和安全边界验证等核心链路。",
    )
    add_mixed_label_para(doc, "关键字：", "大语言模型  Agent运行时  检索增强生成  工业检测软件  DAC-3D")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = Pt(20)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(TITLE_EN)
    set_run_font(r, "Times New Roman", 14, False)

    add_mixed_label_para(
        doc,
        "Abstract: ",
        "DAC-3D inspection software involves parameter setting, scan-region confirmation, runtime status checking, inspection-result interpretation and exception handling. Conventional interfaces can support basic inspection control, but they still have limitations in knowledge retrieval, natural-language operation, result explanation, operation confirmation and process traceability. To address these issues, this paper designs and implements a large-language-model-based intelligent interaction assistant for the DAC-3D scenario. The system adopts a front-end/back-end separated and modular architecture. A unified Agent runtime coordinates knowledge retrieval, context building, skill selection, session memory, tool gateway, safety review, DAC-3D adaptation, result parsing and Trace/Eval modules. The assistant does not change the main detection algorithm or the main inspection workflow of DAC-3D; instead, it builds a knowledge-constrained, tool-gateway-controlled and traceable auxiliary interaction layer outside the main system. Test results show that the prototype can support parameter Q&A, command preview, status query, result explanation and safety-boundary verification in a local validation environment.",
        english=True,
    )
    add_mixed_label_para(
        doc,
        "Keywords: ",
        "Large language model  Agent runtime  Retrieval-augmented generation  Industrial inspection software  DAC-3D",
        english=True,
    )

    add_section(doc, "1.目的意义")
    for text in [
        "随着高精度光学元件在算力中心、光通信以及精密制造等领域中的应用不断增加，检测软件不仅需要完成稳定的图像采集与缺陷识别，还需要为操作人员提供清晰、可追溯、易理解的交互支持。DAC-3D检测软件在实际使用中需要依次完成样品选择、扫描区域确认、参数设置、检测执行、状态查看以及结果判断等工作。对于熟练用户来说，传统菜单和按钮式界面可以完成基本操作；但对于新用户或普通操作人员来说，参数含义、操作前提和结果字段往往不够直观，容易增加学习成本和误操作风险。",
        "大语言模型在自然语言理解、文本组织和任务对话方面表现出较强能力，而检索增强生成方法可以把模型回答约束在给定知识资料之内，减少脱离现场语境的自由发挥。本文把DAC-3D检测场景作为研究对象，并没有把大语言模型用于替代缺陷检测算法，而是将其作为辅助交互层的生成能力来源，使用户能够通过自然语言完成参数咨询、流程确认、状态查询、命令预览和结果解释等任务。该研究的意义在于，将大语言模型能力与工业检测软件的业务边界、安全控制和可验证证据结合起来，探索一种较低侵入、可复查的智能助手实现方式。",
    ]:
        add_para(doc, text)

    add_section(doc, "2.系统设计与实现")
    for text in [
        "系统采用前后端分离和模块化设计。后端以统一Agent运行时为入口，对用户输入进行任务分派，并根据任务类型选择知识库内容、运行状态、技能说明、会话记忆以及安全策略。参数问答和操作指导主要依赖本地知识库检索结果组织回答；操作类请求先形成结构化命令预览，再进入Tool Gateway、PolicyEngine、SafetyGuard、路径白名单、schema校验、风险分类和确认token校验等环节；状态查询和结果解释则通过DAC-3D适配层读取状态文件、最近检测结果或模拟运行数据，并组织成面向用户的说明。",
        "知识库构建部分将DAC-3D操作手册、参数说明、常见问题和缺陷规则整理为可检索资料，经过文本清洗、分块、向量化和索引持久化后，为参数问答、操作指导和结果解释提供依据。系统在回答时尽量返回来源片段、相关字段和规则依据，使用户能够理解结论来自哪里。意图识别和结构化命令生成并不直接执行检测操作，而是把自然语言中的扫描区域、分辨率、模式、离线目录等字段提取出来，形成可校验、可展示、可追踪的命令预览。",
        "DAC-3D主系统适配采用低侵入式思路。主系统已有划痕、点蚀、飞溅等缺陷检测能力，本文工作不包含检测模型本身的训练、改造或部署，而是围绕已有检测输出完成字段规范化和解释链路构建。适配层通过状态文件桥接、命令文件桥接和嵌入式桥接等方式读取运行状态、最近检测结果和命令回执；执行类请求必须经过安全网关和用户确认，不允许大语言模型绕过审批流程直接写入DAC命令文件。前端基于React与Vite实现聊天式交互界面，用于展示回答内容、来源依据、结构化命令、运行状态和确认结果。",
    ]:
        add_para(doc, text)

    add_section(doc, "3.系统测试与运行效果")
    for text in [
        "系统测试围绕功能、接口、异常边界、安全策略、性能、场景化运行和用户体验等方面展开。功能测试覆盖知识库构建、检索问答、自然语言命令预览、状态查询、结果解释、Web交互以及Agent回归场景；安全红队用例重点验证绕过确认、直接写入命令文件、越权路径和Prompt Injection等输入是否会被拒绝。测试结果显示，参数问答、缺陷判定、命令预览和安全边界等链路能够在本地原型环境中形成较为稳定的运行闭环，结果解释字段输出的稳定性仍是后续需要继续改进的重点。",
        "接口测试主要覆盖健康检查、运行状态、同步聊天、流式聊天、命令审批、记忆补丁、Trace读取和评测运行等接口。非模型类接口响应较快，聊天类接口能够完成对真实外部模型或本地Agent验证模型的调用，但完整响应时间会受到外部模型生成速度影响。性能测试结果表明，本地检索、字段解析、路径校验和风险判断所带来的开销相对有限，主要耗时仍集中在模型生成过程。Tool Gateway、PolicyEngine和SafetyGuard虽然增加了执行前校验步骤，但这些步骤属于确定性的本地逻辑，更适宜被视为工业操作安全所必须承担的成本。",
        "用户体验评估采用3项任务与SUS量表相结合的方式开展。5名参与者需要完成参数问答、自然语言命令预览和检测结果解释三类任务，随后填写10题SUS量表。评估结果中，参数问答任务和结果解释任务均由5名参与者完成，自然语言命令预览任务由4名参与者完成，另有1名参与者因未提供分辨率而未完整完成。SUS平均得分为82.0，整体处于较好的可用性区间。该结果说明参与者能够理解系统的主要交互方式，但命令生成过程中的缺失字段提示仍需要进一步增强。",
    ]:
        add_para(doc, text)

    add_section(doc, "4.总结")
    for text in [
        "本文围绕DAC-3D检测软件在知识查找、自然语言操作、结果解释和安全确认方面的实际需求，完成了一套基于大语言模型的智能交互助手原型。系统把大语言模型、检索增强生成、Agent运行时、工具网关、安全审查、状态桥接和结果解析等模块组织到统一链路当中，使用户能够围绕参数、流程、状态、命令和结果开展自然语言交互。相较于普通聊天系统，该原型更加重视知识来源、结构化字段、执行边界和审计记录，能够更好地适应工业检测软件场景中的可控性要求。",
        "本文的主要工作包括四个方面：第一，完成面向DAC-3D应用场景的知识约束型助手框架，使回答内容尽量以本地知识库和结构化结果字段为依据；第二，实现统一Agent运行时和上下文工程模块，使问答、状态、结果、控制、记忆和安全任务能够在同一入口下处理；第三，设计受控工具网关和低侵入式集成方案，避免模型直接控制主系统；第四，补充功能测试、接口测试、异常测试、Agent回归测试、安全红队测试和用户体验评估，形成了可复查的验证材料。",
        "需要指出的是，当前系统仍属于毕业设计阶段的工程原型。知识库资料、真实设备长时间闭环、多用户权限管理、复杂异常回执和结果解释字段稳定性等方面仍需继续完善。后续工作可从正式手册和现场案例补充、真实DAC-3D设备联调、命令生命周期审计、界面分区优化以及Trace/Eval回归资产沉淀等方面继续推进。总体来看，本文验证了在不改变DAC-3D主检测流程边界的前提下，构建一个受知识约束、受工具网关控制并且可追踪验证的LLM助手原型具有可行性。",
    ]:
        add_para(doc, text)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = Pt(18)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("主要参考文献")
    set_run_font(r, "宋体", 10.5, True)

    refs = [
        "[1] Vaswani A, Shazeer N, Parmar N, et al. Attention is all you need[C]//Advances in Neural Information Processing Systems 30. Red Hook: Curran Associates, 2017:5998-6008.",
        "[2] Lewis P, Perez E, Piktus A, et al. Retrieval-augmented generation for knowledge-intensive NLP tasks[C]//Advances in Neural Information Processing Systems 33. Red Hook: Curran Associates, 2020:9459-9474.",
        "[3] Yao S, Zhao J, Yu D, et al. ReAct: Synergizing reasoning and acting in language models[EB/OL]. (2022-10-06)[2026-05-16]. https://arxiv.org/abs/2210.03629.",
        "[4] Asai A, Wu Z, Wang Y, et al. Self-RAG: Learning to retrieve, generate, and critique through self-reflection[C]//Proceedings of the Twelfth International Conference on Learning Representations. Vienna: ICLR, 2024.",
        "[5] Edge D, Trinh H, Cheng N, et al. From local to global: A graph RAG approach to query-focused summarization[EB/OL]. (2024-04-20)[2026-05-24]. https://arxiv.org/abs/2404.16130.",
        "[6] Zheng T, Grosse E H, Morana S, et al. A review of digital assistants in production and logistics: Applications, benefits, and challenges[J]. International Journal of Production Research, 2024, 62(21):8022-8048.",
        "[7] Tiangolo S. FastAPI documentation[EB/OL]. [2026-05-19]. https://fastapi.tiangolo.com/.",
        "[8] Meta Platforms, Inc. React documentation[EB/OL]. [2026-05-19]. https://react.dev/.",
    ]
    for ref in refs:
        add_ref(doc, ref)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                set_cell_margins(cell)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
